
from __future__ import annotations

import os
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .compare import MISMATCH, PERFECT, ROUNDING
from .config import WaterfallConfig

SHEET = "Mismatches"
_MONEY = "#,##0.00;[Red]-#,##0.00;0.00"
_FILL = {MISMATCH: "F8CBAD", ROUNDING: "FFF2CC"}
_STATUS_ORDER = {MISMATCH: 0, ROUNDING: 1}


def _date(value: Any) -> str:
    number = int(value)
    return f"{number // 10000}-{number // 100 % 100:02d}-{number % 100:02d}"


def _number(value: Any) -> float | None:
    return float(value) if isinstance(value, Decimal) else None


def build_report_frame(
    compared: pd.DataFrame, causes: pd.DataFrame, waterfall: WaterfallConfig
) -> pd.DataFrame:
    flagged = compared[compared["status"] != PERFECT].merge(causes, on="row_id", how="left")
    if not flagged.empty:
        flagged = flagged.assign(
            _order=flagged["status"].map(_STATUS_ORDER), _code=flagged["cause_code"]
        ).sort_values(["_order", "_code", "vkorg", "kunnr_ecc", "matnr_ecc", "pricing_date"])

    out = pd.DataFrame(index=range(len(flagged)))
    flagged = flagged.reset_index(drop=True)
    out["Pricing date"] = [_date(v) for v in flagged["pricing_date"]]
    out["Sales org"] = flagged["vkorg"]
    out["Customer (ECC)"] = flagged["kunnr_ecc"]
    out["Customer (S/4)"] = flagged["kunnr_s4"]
    out["Material (ECC)"] = flagged["matnr_ecc"]
    out["Material (S/4)"] = flagged["matnr_s4"]
    out["Status"] = flagged["status"]
    out["Probable cause"] = flagged["probable_cause"]
    out["Confidence"] = flagged["confidence"]
    out["Cause detail"] = flagged["cause_detail"]
    out["Recommended action"] = flagged["recommended_action"]
    out["All cause codes"] = flagged["all_cause_codes"]
    for step in waterfall.procedure:
        out[f"{step.label} ECC"] = [_number(v) for v in flagged[f"ecc_{step.id}"]]
        out[f"{step.label} S/4"] = [_number(v) for v in flagged[f"s4_{step.id}"]]
        out[f"{step.label} Δ"] = [_number(v) for v in flagged[f"delta_{step.id}"]]
    out["Currency ECC"] = flagged["ecc_currency"]
    out["Currency S/4"] = flagged["s4_currency"]
    out["UoM ECC"] = flagged["ecc_uom"]
    out["UoM S/4"] = flagged["s4_uom"]
    out["ECC resolved records"] = flagged["ecc_records"]
    out["S/4 resolved records"] = flagged["s4_records"]
    return out


def write_excel(frame: pd.DataFrame, path: Path, title: str, description: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.parent / f".tmp-{path.name}"
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=SHEET, index=False)
        sheet = writer.sheets[SHEET]

        header_fill = PatternFill("solid", fgColor="1F3864")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32

        columns = list(frame.columns)
        status_col = columns.index("Status") + 1 if "Status" in columns else None
        money_cols = [
            i + 1
            for i, name in enumerate(columns)
            if name.endswith(("ECC", "S/4", "Δ"))
            and not name.startswith(
                ("Customer", "Material", "Currency", "UoM", "ECC resolved", "S/4 resolved")
            )
        ]
        delta_cols = [i + 1 for i, name in enumerate(columns) if name.endswith("Δ")]
        wide = {
            "Cause detail": 70,
            "Recommended action": 55,
            "ECC resolved records": 60,
            "S/4 resolved records": 60,
            "Probable cause": 42,
        }

        for row in sheet.iter_rows(min_row=2):
            status = row[status_col - 1].value if status_col else None
            if status in _FILL:
                row[status_col - 1].fill = PatternFill("solid", fgColor=_FILL[status])
            for col in money_cols:
                row[col - 1].number_format = _MONEY
            for col in delta_cols:
                cell = row[col - 1]
                if cell.value not in (None, 0):
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="FCE4D6")

        for index, name in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = wide.get(
                name, max(12, min(18, len(name) + 2))
            )

        sheet.freeze_panes = "G2"
        if len(frame):
            sheet.auto_filter.ref = sheet.dimensions

        workbook = writer.book
        workbook.properties.title = title
        workbook.properties.description = description
        workbook.properties.creator = "pricing-migration-reconciler"
        workbook.properties.created = datetime.now()

    os.replace(temporary, path)
    return path
